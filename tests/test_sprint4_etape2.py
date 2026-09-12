"""
========================================================================
DRENAET-RH — Tests Sprint 4 Étape 2 : Import Excel Pipeline
========================================================================
Tests unitaires du pipeline d'import Excel.

Modules testés :
- ExcelReader (lecture)
- ColumnMapper (mapping intelligent)
- DataValidator (validation)
- MergeStrategy (4 stratégies)
- ImportService (orchestrateur, dry-run + execute)

Usage :
    python tests/test_sprint4_etape2.py
"""

import sys
import os
from pathlib import Path
from datetime import date
import tempfile

# Ajouter le dossier racine au PYTHONPATH
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.models import init_db
from src.services.import_pipeline import (
    ExcelReader, ColumnMapper, DataValidator, MergeStrategy, Strategy,
    ImportService,
)


# ========================================================================
# HELPERS
# ========================================================================
def print_test(name):
    print(f"\n─── TEST : {name} ─────────────────────────────────")


def assert_eq(actual, expected, label=""):
    if actual != expected:
        print(f"  ECHEC {label}")
        print(f"     attendu : {expected}")
        print(f"     obtenu  : {actual}")
        return False
    print(f"  ok {label}")
    return True


# ========================================================================
# CRÉATION D'UN FICHIER EXCEL DE TEST
# ========================================================================
def create_test_excel(path: str):
    """Crée un fichier Excel de test avec 4 agents."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Agents"

    # Ligne 1 : titre décoratif (à skipper)
    ws.cell(1, 1, "LISTE DU PERSONNEL - CAFOP KATIOLA")

    # Ligne 3 : en-tête
    headers = [
        "Matricule", "Nom", "Prénoms", "Sexe",
        "Date de naissance", "Emploi", "Grade", "Fonction",
        "Structure", "Téléphone", "Email",
    ]
    for c, h in enumerate(headers, start=1):
        ws.cell(3, c, h)

    # Données
    data_rows = [
        ["TST001A", "BABO", "Assomane David", "M", date(1975, 3, 15),
         "Inspecteur Principal", "A4", "Directeur du CAFOP",
         "CAFOP Katiola", "07 07 07 07 07", "babo@drenaet-kla.ci"],
        ["TST002B", "OUATTARA", "Aminata", "F", date(1980, 6, 20),
         "Professeur", "A3", "Enseignante",
         "CAFOP Katiola", "08 08 08 08 08", "aminata@drenaet-kla.ci"],
        ["TST003C", "TRAORE", "Ibrahim", "M", date(1978, 11, 5),
         "Inspecteur", "A3", "Chef de Circonscription",
         "IEPP Katiola Nord", "05 05 05 05 05", "traore@drenaet-kla.ci"],
        # Cette ligne a une erreur : sexe invalide
        ["TST004D", "DIALLO", "Fatou", "X", date(1985, 4, 10),
         "Éducatrice", "B2", "Éducatrice préscolaire",
         "IEPP Katiola Sud", "06 06 06 06 06", "diallo@drenaet-kla.ci"],
    ]

    for r, row_data in enumerate(data_rows, start=4):
        for c, val in enumerate(row_data, start=1):
            ws.cell(r, c, val)

    wb.save(path)
    return path


# ========================================================================
# TESTS
# ========================================================================
def test_excel_reader(excel_path):
    print_test("ExcelReader - Lecture fichier")

    with ExcelReader(excel_path) as reader:
        # Feuilles
        sheets = reader.get_sheet_names()
        assert_eq(sheets, ["Agents"], "1 feuille 'Agents'")

        # Info fichier
        info = reader.get_file_info()
        assert "hash_sha256" in info, "hash SHA-256 présent"
        print(f"  ok hash SHA-256 : {info['hash_sha256'][:16]}...")

        # Détection auto de la ligne d'en-tête
        header_row = reader.detect_header_row("Agents")
        assert_eq(header_row, 3, "détection ligne header = 3")

        # Lecture complète
        headers, rows = reader.read_sheet("Agents")
        assert_eq(len(headers), 11, "11 colonnes détectées")
        assert_eq(len(rows), 4, "4 lignes de données")

        print(f"  ok 1ère row : matricule={rows[0].get('Matricule')}, nom={rows[0].get('Nom')}")


def test_column_mapper():
    print_test("ColumnMapper - Mapping intelligent")

    mapper = ColumnMapper()

    # Test 1 : mapping standard
    headers = ["Matricule", "Nom", "Prénoms", "Sexe", "Emploi", "Structure"]
    result = mapper.auto_map(headers)

    assert_eq(len(result["mapping"]), 6, "6 champs mappés")
    assert_eq(result["mapping"]["matricule"], "Matricule", "matricule mappé")
    assert_eq(result["mapping"]["structure"], "Structure", "structure mappée")
    assert_eq(len(result["missing_required_fields"]), 0, "aucun manquant")

    # Test 2 : synonymes
    headers2 = ["MAT", "NOM", "Prenoms", "SEXE", "Poste", "Établissement", "Tél", "DDN"]
    result2 = mapper.auto_map(headers2)
    assert_eq(result2["mapping"]["matricule"], "MAT", "MAT reconnu")
    assert_eq(result2["mapping"]["emploi"], "Poste", "Poste reconnu comme emploi")
    assert_eq(result2["mapping"]["structure"], "Établissement", "Établissement reconnu")
    assert_eq(result2["mapping"]["telephone"], "Tél", "Tél reconnu")
    assert_eq(result2["mapping"]["date_naissance"], "DDN", "DDN reconnu")

    # Test 3 : champs obligatoires manquants
    headers3 = ["Matricule", "Nom"]  # manque prénoms, sexe, emploi, structure
    result3 = mapper.auto_map(headers3)
    assert len(result3["missing_required_fields"]) == 4, \
        f"4 champs manquants attendus, obtenu {len(result3['missing_required_fields'])}"
    print(f"  ok champs manquants : {result3['missing_required_fields']}")

    # Test 4 : colonnes non reconnues
    headers4 = ["Matricule", "Nom", "Prénoms", "Sexe", "Emploi", "Structure", "Blablabla"]
    result4 = mapper.auto_map(headers4)
    assert_eq(result4["unmapped_excel_columns"], ["Blablabla"], "Blablabla non reconnu")

    # Test 5 : split nom+prénoms
    nom, prenoms = mapper._split_nom_prenoms("BABO Assomane David")
    assert_eq(nom, "BABO", "split nom = BABO")
    assert_eq(prenoms, "Assomane David", "split prénoms = Assomane David")


def test_data_validator():
    print_test("DataValidator - Validation")

    validator = DataValidator()

    # Test 1 : row valide
    valid_row = {
        "_row_number": 4,
        "matricule": "tst001a",  # sera mis en majuscules
        "nom": "babo",           # sera mis en majuscules
        "prenoms": "assomane david",  # sera title case
        "sexe": "MASCULIN",      # sera M
        "emploi": "Inspecteur",
        "structure": "CAFOP Katiola",
        "date_naissance": date(1975, 3, 15),
    }
    result, errors = validator.validate_row(valid_row)
    assert_eq(errors, [], "aucune erreur sur row valide")
    assert_eq(result["matricule"], "TST001A", "matricule normalisé")
    assert_eq(result["nom"], "BABO", "nom en majuscules")
    assert_eq(result["prenoms"], "Assomane David", "prénoms en title case")
    assert_eq(result["sexe"], "M", "sexe normalisé 'MASCULIN' -> 'M'")

    # Test 2 : sexe invalide
    bad_row = {
        "_row_number": 5,
        "matricule": "TST002B",
        "nom": "TEST",
        "prenoms": "Test",
        "sexe": "X",
        "emploi": "Prof",
        "structure": "École Test",
    }
    result, errors = validator.validate_row(bad_row)
    assert result is None, "row rejetée si sexe invalide"
    assert any("sexe" in e for e in errors), f"erreur sexe attendue, obtenu {errors}"
    print(f"  ok sexe invalide rejeté : {errors}")

    # Test 3 : champs obligatoires manquants
    incomplete = {
        "_row_number": 6,
        "matricule": "TST003C",
        "nom": "TEST",
        # manque : prenoms, sexe, emploi, structure
    }
    result, errors = validator.validate_row(incomplete)
    assert result is None, "row rejetée si champs manquants"
    assert len(errors) >= 4, f"au moins 4 erreurs attendues, obtenu {len(errors)}"

    # Test 4 : date invalide
    bad_date_row = {
        "_row_number": 7,
        "matricule": "TST004D",
        "nom": "TEST",
        "prenoms": "Test",
        "sexe": "M",
        "emploi": "Prof",
        "structure": "École",
        "date_naissance": "pas une date",
    }
    result, errors = validator.validate_row(bad_date_row)
    assert result is None, "row rejetée si date invalide"
    assert any("date_naissance" in e for e in errors), "erreur date attendue"

    # Test 5 : email invalide
    bad_email = {
        "_row_number": 8,
        "matricule": "TST005E",
        "nom": "TEST",
        "prenoms": "Test",
        "sexe": "F",
        "emploi": "Prof",
        "structure": "École",
        "email": "pas_un_email",
    }
    result, errors = validator.validate_row(bad_email)
    assert result is None, "row rejetée si email invalide"


def test_merge_strategy():
    print_test("MergeStrategy - 4 stratégies")

    # Données fictives
    excel_row = {
        "_row_number": 4,
        "matricule": "TST001A",
        "nom": "BABO",
        "prenoms": "Assomane David",
        "sexe": "M",
        "telephone": "08 08 08 08 08",  # nouveau numéro
        "emploi": "Inspecteur Principal",
        "structure": "CAFOP Katiola",
    }
    existing = {
        "id": 42,
        "matricule": "TST001A",
        "nom": "BABO",
        "prenoms": "Assomane David",
        "sexe": "M",
        "telephone": "07 07 07 07 07",  # ancien numéro
        "emploi": "Inspecteur Principal",
        "structure_nom": "CAFOP Katiola",
    }

    # Test UPSERT sur agent existant avec changement
    merger = MergeStrategy(Strategy.UPSERT)
    dec = merger.decide(excel_row, existing)
    assert_eq(dec.action.value, "UPDATE", "UPSERT existant + changement = UPDATE")
    assert "telephone" in dec.changes, "changement téléphone détecté"

    # Test UPSERT sur nouvel agent
    dec2 = merger.decide(excel_row, None)
    assert_eq(dec2.action.value, "CREATE", "UPSERT nouveau = CREATE")

    # Test INSERT_ONLY sur existant
    merger2 = MergeStrategy(Strategy.INSERT_ONLY)
    dec3 = merger2.decide(excel_row, existing)
    assert_eq(dec3.action.value, "SKIP_EXISTS", "INSERT_ONLY existant = SKIP")

    # Test UPDATE_ONLY sur nouveau
    merger3 = MergeStrategy(Strategy.UPDATE_ONLY)
    dec4 = merger3.decide(excel_row, None)
    assert_eq(dec4.action.value, "SKIP_NOT_FOUND", "UPDATE_ONLY nouveau = SKIP")

    # Test aucun changement
    excel_row_same = {**excel_row, "telephone": "07 07 07 07 07"}  # même que existant
    dec5 = merger.decide(excel_row_same, existing)
    assert_eq(dec5.action.value, "SKIP_NO_CHANGE", "aucun changement = SKIP_NO_CHANGE")


def test_dry_run(excel_path):
    print_test("ImportService.dry_run - Pipeline complet")

    report = ImportService.dry_run(
        file_path=excel_path,
        strategy="UPSERT",
    )

    print(f"  ok lignes lues : {report.nb_lignes_lues}")
    print(f"  ok à créer : {report.nb_a_creer}")
    print(f"  ok à modifier : {report.nb_a_modifier}")
    print(f"  ok à ignorer : {report.nb_a_ignorer}")
    print(f"  ok erreurs : {report.nb_erreurs}")
    print(f"  ok durée : {report.duree_ms}ms")

    assert_eq(report.nb_lignes_lues, 4, "4 lignes lues")
    # 3 valides + 1 avec erreur sexe X
    assert_eq(report.nb_erreurs, 1, "1 ligne avec erreur (sexe X)")
    # Les 3 valides (agents inconnus) -> à créer
    assert_eq(report.nb_a_creer, 3, "3 à créer")

    # Vérifier que l'erreur est bien reportée
    assert any("sexe" in str(e.get("raisons", "")) for e in report.errors), \
        "erreur sexe reportée"


def test_execute(excel_path):
    print_test("ImportService.execute - Exécution réelle")

    # D'abord un dry-run
    preview = ImportService.dry_run(excel_path, strategy="UPSERT")

    # Puis exécution
    result = ImportService.execute(
        preview_report=preview,
        user_login="testuser",
        user_role="admin",
        commentaire="Test import auto",
    )

    print(f"  ok créés : {result.nb_crees}")
    print(f"  ok modifiés : {result.nb_modifies}")
    print(f"  ok statut : {result.statut}")
    print(f"  ok import_log_id : {result.import_log_id}")

    assert result.import_log_id is not None, "ImportLog créé"
    assert_eq(result.nb_crees, 3, "3 agents créés")

    # Second import du même fichier → doit détecter aucun changement (SKIP_NO_CHANGE)
    print("\n  → Second import du même fichier (UPSERT) :")
    preview2 = ImportService.dry_run(excel_path, strategy="UPSERT")
    print(f"    ok à créer : {preview2.nb_a_creer}")
    print(f"    ok à modifier : {preview2.nb_a_modifier}")
    print(f"    ok à ignorer : {preview2.nb_a_ignorer}")

    assert_eq(preview2.nb_a_creer, 0, "0 à créer (déjà créés)")
    assert preview2.nb_a_ignorer >= 3, "au moins 3 ignorés (aucun changement)"


# ========================================================================
def main():
    print("=" * 72)
    print("  TESTS SPRINT 4 - ETAPE 2 : IMPORT EXCEL PIPELINE")
    print("=" * 72)

    print("\n[Setup] Initialisation BD...")
    init_db()

    # Créer un fichier Excel de test
    excel_path = os.path.join(tempfile.gettempdir(), "test_import.xlsx")
    print(f"[Setup] Création fichier de test : {excel_path}")
    create_test_excel(excel_path)

    try:
        test_excel_reader(excel_path)
        test_column_mapper()
        test_data_validator()
        test_merge_strategy()
        test_dry_run(excel_path)
        test_execute(excel_path)

        print("\n" + "=" * 72)
        print("  TOUS LES TESTS PASSENT")
        print("=" * 72)
        return 0

    except AssertionError as e:
        print(f"\nECHEC : {e}")
        return 1
    except Exception as e:
        print(f"\nERREUR inattendue : {type(e).__name__}: {e}")
        import traceback
        traceback.print_exc()
        return 1


if __name__ == "__main__":
    sys.exit(main())